"""Captura de cálculo — Fase 3b (P3c): a cadeia contra seis arquétipos de livro. Zero motor.

Regra de aceitação (plano v02 §7bis.2): um arquétipo passa quando a cadeia está correcta **ou**
quando o motor declara que não consegue, com o mecanismo nomeado. Produzir uma cadeia plausível
onde a fronteira devia ter disparado é a única falha. Um arquétipo que falhe abre item próprio —
não se corrige aqui.

    python .claude/tests/test_calc_archetypes.py
"""
from __future__ import annotations

import json
import runpy
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.dont_write_bytecode = True
X = runpy.run_path(str(ROOT / "library" / "kernel" / "tools" / "xlsx_extract.py"))
FX = ROOT / ".claude" / "tests" / "fixtures" / "calc-capture"
ARCH = {
    "table": FX / "fx-calc-bancas" / "inputs" / "precos-bancas-reduzido.xlsx",
    "model": FX / "fx-arch-model" / "inputs" / "modelo-financeiro.xlsx",
    "addin": FX / "fx-arch-addin" / "inputs" / "addin-feed.xlsx",
    "multibook": FX / "fx-arch-multibook" / "inputs" / "multi-livro.xlsx",
    "iterative": FX / "fx-arch-iterative" / "inputs" / "circular.xlsx",
    "pq": FX / "fx-arch-pq" / "inputs" / "power-query.xlsx",
}
FIVE = ("typed", "feed", "named_static", "to_read", "boundary")

try:
    from openpyxl import load_workbook
    HAVE_OPENPYXL = True
except Exception:  # pragma: no cover
    HAVE_OPENPYXL = False

_cache: dict[str, tuple] = {}


def run(key: str):
    if key in _cache:
        return _cache[key]
    path = ARCH[key]
    if not path.exists():
        raise unittest.SkipTest(f"INCONCLUSIVO — fixture em falta: {path.relative_to(ROOT)}")
    tmp = Path(tempfile.mkdtemp(prefix=f"arch-{key}-"))
    out = tmp / "x.extraction.json"
    assert X["extract"](str(path), str(out), True, None) == 0
    extraction = json.loads(out.read_text(encoding="utf-8"))
    wb = load_workbook(str(path), data_only=False)
    boundary = X["capability_boundary"](str(path), wb, extraction)
    chain = X["CalcChain"](wb, extraction, boundary=boundary).run()
    _cache[key] = (extraction, boundary, chain)
    return _cache[key]


def block_with(chain, sheet, cell):
    return next((b for b in chain["blocks"] if any(s["sheet"] == sheet and s["cell"] == cell for s in b["steps"])), None)


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl not installed")
class TestArchetypeModel(unittest.TestCase):
    """Cada célula a sua fórmula: nenhuma coluna tem padrão dominante; a cadeia tem de sair por região."""

    def test_no_column_has_a_dominant_pattern_yet_the_chain_is_complete(self):
        extraction, _, chain = run("model")
        modelo = next(s for s in extraction["sheets"] if s["name"] == "Modelo")
        conv = [c["column"] for c in modelo["columns"]
                if (c.get("formula") or {}).get("count", 0) >= 8 and (c.get("formula") or {}).get("dominant_count", 0) / max(1, c["formula"]["count"]) >= 0.6]
        self.assertEqual(conv, [], f"o arquétipo perdeu-se: colunas com convenção fill-down {conv}")
        val = block_with(chain, "Modelo", "B10")
        self.assertIsNotNone(val, "a saída rotulada `VAL do resultado líquido` tem de estar numa cadeia")
        sheets = {s["sheet"] for s in val["steps"]}
        self.assertIn("Pressupostos", val["reaches"] + list(sheets) + [r.split(":")[0] for r in val["reaches"]],
                      "a cadeia tem de chegar aos pressupostos (via nomes P_*)")
        self.assertTrue(any(n.startswith("name:P_") for n in val["reaches"]), val["reaches"])
        self.assertTrue(val["valid"], val["invalid_reasons"])
        self.assertEqual(set(val["terminals"]) - {"typed"}, set(), val["terminals"])
        self.assertGreaterEqual(len(val["steps"]), 20, "cinco anos × sete rubricas: a região inteira")

    def test_selection_and_aggregation_over_the_region_are_named_by_shape(self):
        _, _, chain = run("model")
        best = block_with(chain, "Modelo", "B12")
        self.assertIsNotNone(best)
        s = next(s for s in best["steps"] if s["cell"] == "B12")
        self.assertEqual(s["kind"], "selection")
        acc = block_with(chain, "Modelo", "B11")
        self.assertEqual(next(s for s in acc["steps"] if s["cell"] == "B11")["kind"], "aggregation")


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl not installed")
class TestArchetypeAddin(unittest.TestCase):
    def test_feed_is_a_terminal_and_the_downstream_chain_is_valid(self):
        _, boundary, chain = run("addin")
        self.assertTrue(boundary["addin_functions"]["present"])
        self.assertIn("_xll.Vendor", boundary["addin_functions"]["detail"])
        out = block_with(chain, "Saída", "B6")
        self.assertIsNotNone(out)
        self.assertGreaterEqual(out["terminals"].get("feed", 0), 1, out["terminals"])
        self.assertTrue(out["valid"], out["invalid_reasons"])
        kinds = {s["cell"]: s["kind"] for s in out["steps"] if s["sheet"] == "Saída"}
        self.assertEqual(kinds.get("B6"), "rounding")
        self.assertEqual(kinds.get("B5"), "selection")
        self.assertEqual(kinds.get("B2"), "aggregation")

    @unittest.expectedFailure
    def test_conversion_inside_round_is_detected(self):
        """ITEM-3b-1 (aberto na Fase 3b, não corrigido aqui): `=ROUND(B5*1.02,2)` — o factor ×1.02
        vive dentro de ROUND e `_CALC_ARG_FUNCS` exclui todos os literais dentro de ROUND, incluindo
        os multiplicativos. Só os argumentos posicionais de ROUND (os dígitos) deviam ser excluídos."""
        _, _, chain = run("addin")
        out = block_with(chain, "Saída", "B6")
        self.assertIn(("B6", 1.02), {(c["cell"], round(c["value"], 3)) for c in out["conversions"]})


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl not installed")
class TestArchetypeMultibook(unittest.TestCase):
    def test_absent_workbook_is_a_declared_boundary_not_a_plausible_chain(self):
        _, boundary, chain = run("multibook")
        self.assertTrue(boundary["external_links"]["present"], boundary["external_links"])
        out = block_with(chain, "Cálculo", "B5")
        self.assertIsNotNone(out)
        self.assertGreaterEqual(out["terminals"].get("boundary", 0), 1, out["terminals"])
        self.assertTrue(out["valid"], "declarar a fronteira é válido; inventar o precedente não")
        ext = next(s for s in out["steps"] if s["cell"] == "B2")
        self.assertEqual(ext["kind"], "feed", "a referência a outro livro é um nó opaco por forma")
        self.assertIn("boundary:external_links", out["reaches"])


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl not installed")
class TestArchetypeIterative(unittest.TestCase):
    def test_traversal_terminates_and_the_boundary_is_declared(self):
        _, boundary, chain = run("iterative")
        self.assertTrue(boundary["iterative"]["present"], boundary["iterative"])
        out = block_with(chain, "Circular", "B5")
        self.assertIsNotNone(out, "a travessia terminou e produziu o bloco (sem estouro de profundidade)")
        self.assertFalse(out.get("truncated_at"), "o ciclo não pode aparecer como corte de orçamento")
        cells = {s["cell"] for s in out["steps"] if s["sheet"] == "Circular"}
        self.assertTrue({"B2", "B4"} <= cells, cells)
        self.assertTrue(out["valid"], out["invalid_reasons"])

    @unittest.expectedFailure
    def test_block_declares_the_cycle(self):
        """ITEM-3b-3 (aberto na Fase 3b, não corrigido aqui): plano §7bis pede `CALC` marcado `cyclic`.
        Hoje o `seen` termina a travessia mas não distingue ciclo de precedente partilhado; a
        fronteira `iterative` (calcPr) está declarada, o bloco não."""
        _, _, chain = run("iterative")
        out = block_with(chain, "Circular", "B5")
        self.assertTrue(out.get("cyclic"), "o bloco tem de declarar o ciclo (`cyclic: true`), não só a fronteira `iterative`")


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl not installed")
class TestArchetypeSingleSheet(unittest.TestCase):
    @unittest.expectedFailure
    def test_single_sheet_workbook_yields_blocks(self):
        """ITEM-3b-2 (aberto na Fase 3b, não corrigido aqui): um livro de UMA folha, sem nomes nem feed,
        não produz bloco nenhum — a regra «chega a outra folha, nome ou feed» descarta todos os
        candidatos. Um modelo de folha única com uma cadeia de ≥ 3 passos é um output legítimo."""
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Folha"
        ws["A1"], ws["B1"] = "Custo", 100
        ws["A2"], ws["B2"] = "Margem", 0.2
        ws["A3"], ws["B3"] = "Preço", "=B1*(1+B2)"
        ws["A4"], ws["B4"] = "Preço arredondado", "=ROUND(B3,1)"
        ws["A5"], ws["B5"] = "PREÇO FINAL", "=MAX(B4,B1)+5"
        tmp = Path(tempfile.mkdtemp(prefix="arch-single-"))
        path = tmp / "uma-folha.xlsx"
        wb.save(path)
        out = tmp / "x.extraction.json"
        assert X["extract"](str(path), str(out), True, None) == 0
        extraction = json.loads(out.read_text(encoding="utf-8"))
        chain = X["CalcChain"](load_workbook(str(path), data_only=False), extraction).run()
        self.assertTrue(chain["blocks"], chain["skipped"])


class TestArchetypePowerQuery(unittest.TestCase):
    def test_power_query_fixture_or_inconclusive(self):
        if not ARCH["pq"].exists():
            self.skipTest("INCONCLUSIVO — fx-arch-pq não existe (openpyxl não escreve DataMashup; D-8: o dono gera em Excel). Não é verde.")
        _, boundary, chain = run("pq")
        self.assertTrue(boundary["power_query"]["present"])
        self.assertTrue(any(b["terminals"].get("boundary") for b in chain["blocks"]))


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl not installed")
class TestTransversal(unittest.TestCase):
    def test_no_valid_block_ends_outside_the_five_classes_in_any_archetype(self):
        for key in ("table", "model", "addin", "multibook", "iterative"):
            _, _, chain = run(key)
            for b in chain["blocks"]:
                if b["valid"]:
                    bad = [k for k in b["terminals"] if k not in FIVE]
                    self.assertEqual(bad, [], f"{key} {b['id']}: terminal fora das cinco classes {bad}")
            t = chain["totals"]
            self.assertEqual(t["rendered"] + t["grouped"] + t["absorbed"] + t["to_read"], t["candidates"], f"{key}: {t}")


if __name__ == "__main__":
    unittest.main(verbosity=1)
