#!/usr/bin/env python3
"""Gera os dois workbooks sinteticos da fixture `fx-coverage-f06` e corre os motores
reais de captura sobre eles (xlsx_extract L1+L3, fields_draft L1b).

    python .claude/tests/fixtures/coverage/build_workbooks.py <repo-root>

Nada aqui e especifico de nenhum caso real: o dominio e laboratorial, inventado para o
efeito. O que se preserva do caso F06 e a FORMA -- duas folhas de saida, uma enviada por
email e outra apenas gravada para consulta, mais um grupo de colunas repetidas e um
segundo workbook com uma folha do mesmo nome.
"""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook

# Mesma regra do `build_fixtures.py`: a raiz vem de `__file__`, nunca de `sys.argv`.
REPO = Path(__file__).resolve().parents[4]
FIX = REPO / ".claude" / "tests" / "fixtures" / "coverage" / "fx-coverage-f06"
TOOLS = REPO / "library" / "kernel" / "tools"
INPUTS = FIX / "inputs"
CAPTURE = FIX / "_capture"

PRIMARY = "registo-de-lotes.xlsx"
SECOND = "registo-de-lotes-2024.xlsx"


def build_primary(path: Path) -> None:
    wb = Workbook()

    ws = wb.active
    ws.title = "Entradas"
    ws.append(["Data", "Lote", "Produto", "ValorMedido", "Operador",
               "Custo_Posto_A", "Custo_Posto_B", "Custo_Posto_C"])
    for i in range(1, 13):
        ws.append(["2026-01-%02d" % i, "L%03d" % i, "P-1" if i % 2 else "P-2",
                   10 + i * 0.5, "op%d" % (i % 3), 4.0 + i, 4.5 + i, 5.0 + i])

    ws = wb.create_sheet("Resumo de Lotes")
    ws.append(["Data", "Lote", "Resultado", "Estado"])
    for i in range(1, 13):
        r = i + 1
        ws.append(["=Entradas!A%d" % r, "=Entradas!B%d" % r,
                   "=Entradas!D%d*Parametros!$B$2", "Libertado"])

    ws = wb.create_sheet("Resumo Aditivos")
    ws.append(["Data", "Lote", "Aditivo", "Percentagem"])
    for i in range(1, 13):
        r = i + 1
        ws.append(["=Entradas!A%d" % r, "=Entradas!B%d" % r, "AD-1",
                   "=Entradas!D%d/Parametros!$B$3" % r])

    ws = wb.create_sheet("Parametros")
    ws.append(["Chave", "Valor"])
    ws.append(["factor_resultado", 1.2])
    ws.append(["base_aditivo", 40])
    ws.append(["limite_alerta", 95])

    wb.defined_names.add(_dn("Factor_Resultado", "Parametros!$B$2"))
    wb.defined_names.add(_dn("Base_Aditivo", "Parametros!$B$3"))
    wb.defined_names.add(_dn("Lote_Lista", "Entradas!$B$2:$B$13"))
    # Entrada de dicionario sem dados: o nome existe, a coluna esta vazia.
    wb.defined_names.add(_dn("Custo_Posto_D", "Entradas!$I$2:$I$13"))
    wb.save(path)


def build_second(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Entradas"
    ws.append(["Data", "Lote", "ValorMedido"])
    for i in range(1, 6):
        ws.append(["2024-06-%02d" % i, "H%03d" % i, 9 + i])
    # MESMO nome de folha do workbook primario: as unidades nao podem colidir.
    ws = wb.create_sheet("Resumo Aditivos")
    ws.append(["Data", "Lote", "Aditivo", "Percentagem"])
    for i in range(1, 6):
        ws.append(["2024-06-%02d" % i, "H%03d" % i, "AD-0", 0.1 * i])
    wb.save(path)


def _dn(name: str, ref: str):
    from openpyxl.workbook.defined_name import DefinedName
    return DefinedName(name, attr_text="'" + ref.split("!")[0] + "'!" + ref.split("!")[1]
                       if " " in ref.split("!")[0] else ref)


def run(args: list[str]) -> None:
    print("$", " ".join(str(a) for a in args))
    r = subprocess.run([sys.executable, "-B"] + [str(a) for a in args],
                       capture_output=True, text=True)
    sys.stdout.write(r.stdout)
    sys.stderr.write(r.stderr)
    if r.returncode not in (0,):
        raise SystemExit("motor de captura falhou: rc=%d" % r.returncode)


def main() -> None:
    if len(sys.argv) > 1 and Path(sys.argv[1]).resolve() != REPO:
        raise SystemExit(
            "raiz do repositorio nao corresponde: argumento {} != {} (derivada de "
            "__file__).".format(Path(sys.argv[1]).resolve(), REPO))
    INPUTS.mkdir(parents=True, exist_ok=True)
    CAPTURE.mkdir(parents=True, exist_ok=True)

    build_primary(INPUTS / PRIMARY)
    build_second(INPUTS / SECOND)
    # Formato fora dos tiers de captura: fica no manifesto como `not captured`.
    (INPUTS / "fluxo-de-libertacao.pptx").write_bytes(
        b"FIXTURE PLACEHOLDER - formato fora dos tiers de captura, sem conteudo real.\n")

    log = CAPTURE / "_capture-log.md"
    if not log.exists():
        log.write_text("# Capture log\n\n| timestamp | layer | file | event | detail |\n"
                       "|---|---|---|---|---|\n", encoding="utf-8")

    for wbname in (PRIMARY, SECOND):
        run([TOOLS / "xlsx_extract.py", INPUTS / wbname,
             CAPTURE / (wbname + ".extraction.json"), "--force", "--log", log])
    run([TOOLS / "xlsx_extract.py", "--replay", INPUTS / PRIMARY,
         CAPTURE / (PRIMARY + ".extraction.json"),
         CAPTURE / (PRIMARY + ".replay.md"), "--log", log])
    run([TOOLS / "fields_draft.py", CAPTURE / (PRIMARY + ".extraction.json"),
         CAPTURE / (PRIMARY + ".fields-draft.json"), "--force", "--log", log])
    print("ok")


if __name__ == "__main__":
    main()
