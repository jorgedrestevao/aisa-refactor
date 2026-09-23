"""Captura de cálculo — Fase 3c: os limites do extractor são declarados, e a falha degrada até ao fim.

Medido no workbook real do dono (`Preços Transportes`, 13,3 MB): uma folha declara 1 048 576 linhas
para 1,1 M células — colunas formatadas por inteiro. `scan_columns` percorria `data_start..ws.max_row`
com `ws.cell(r, c)`, que **cria** a célula quando não existe: pico de **16,3 GB** e 357 s, e
`MemoryError` quando a máquina não tinha folga. O tecto `MAX_SCAN_ROWS` já existia e só era aplicado
ao varrimento de preenchimentos.

O que este teste fixa:
  1. o varrimento pára no tecto, e a truncagem é **declarada** (`data_end_declared`, `scan_truncated`);
  2. `data_end_row` é o valor efectivamente varrido — o replay herda o tecto por o ler;
  3. os acumuladores por coluna têm tecto, e o que foi contado além dele é dito;
  4. um `MemoryError` produz artefacto de falha **e** termina com rc=0 — a degradação fecha.

    python .claude/tests/test_scan_limits.py
"""
from __future__ import annotations

import json
import runpy
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.dont_write_bytecode = True
X = runpy.run_path(str(ROOT / "library" / "kernel" / "tools" / "xlsx_extract.py"))
CAP = X["MAX_SCAN_ROWS"]

try:
    from openpyxl import Workbook, load_workbook
    HAVE_OPENPYXL = True
except Exception:  # pragma: no cover
    HAVE_OPENPYXL = False

FAR_ROW = CAP * 3          # muito acima do tecto, sem gerar um ficheiro grande
DISTINCT = 500


def wide_workbook(path: Path) -> None:
    """Folha com dados reais no topo e uma célula solitária muito abaixo — a forma que o Excel
    grava quando alguém formata uma coluna inteira."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Larga"
    ws["A1"], ws["B1"], ws["C1"] = "Chave", "Valor", "Derivado"
    for r in range(2, DISTINCT + 2):
        ws[f"A{r}"] = f"K{r:05d}"
        ws[f"B{r}"] = r * 1.5
        ws[f"C{r}"] = f"=B{r}*2"
    ws[f"A{FAR_ROW}"] = "cauda"          # empurra ws.max_row para muito longe
    wb.save(path)


def extract(path: Path) -> dict:
    out = path.with_suffix(".extraction.json")
    rc = X["extract"](str(path), str(out), True, None)
    assert rc == 0, f"extract rc={rc}"
    return json.loads(out.read_text(encoding="utf-8"))


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl not installed")
class TestScanCap(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp(prefix="scan-limits-"))
        cls.path = cls.tmp / "larga.xlsx"
        wide_workbook(cls.path)
        cls.ex = extract(cls.path)
        cls.sheet = next(s for s in cls.ex["sheets"] if s["name"] == "Larga")

    def test_scan_stops_at_the_cap(self):
        s = self.sheet
        self.assertLessEqual(s["data_end_row"] - s["data_start_row"] + 1, CAP,
                             f"varreu {s['data_end_row'] - s['data_start_row'] + 1} linhas, tecto {CAP}")

    def test_truncation_is_declared_not_silent(self):
        s = self.sheet
        self.assertTrue(s.get("scan_truncated"), "a truncagem tem de ser declarada no registo da folha")
        self.assertEqual(s.get("data_end_declared"), FAR_ROW,
                         "o valor declarado pelo ficheiro fica registado ao lado do efectivo")
        self.assertLess(s["data_end_row"], s["data_end_declared"])

    def test_untruncated_sheet_says_so(self):
        small = self.tmp / "pequena.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"], ws["A2"], ws["B2"] = "Chave", "K1", 10
        wb.save(small)
        sheet = extract(small)["sheets"][0]
        self.assertFalse(sheet.get("scan_truncated"), sheet)
        self.assertEqual(sheet["data_end_row"], sheet.get("data_end_declared"))

    def test_replay_inherits_the_cap_from_the_record(self):
        """Os três loops do replay lêem `data_end_row`; com o valor efectivo, herdam o tecto."""
        rep = X["Replayer"](load_workbook(str(self.path), data_only=True), self.ex)
        rep._fwb = load_workbook(str(self.path), data_only=False)
        rep.run()
        for f in rep.findings:
            rows = [int(x) for x in __import__("re").findall(r"\d+", f.location)]
            self.assertTrue(all(r <= self.sheet["data_end_row"] or r < 10 for r in rows),
                            f"finding fora do intervalo varrido: {f.location}")


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl not installed")
class TestAccumulatorCaps(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp(prefix="scan-acc-"))
        cls.path = cls.tmp / "distintos.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Chave"
        for r in range(2, DISTINCT + 2):
            ws[f"A{r}"] = f"V{r:06d}"
        wb.save(cls.path)
        cls.col = next(c for c in extract(cls.path)["sheets"][0]["columns"] if c["column"] == "A")

    def test_distinct_is_capped_and_the_cap_is_declared(self):
        self.assertIn("MAX_DISTINCT", X, "Fase 3c: tecto declarado para o contador de valores distintos")
        cap = X["MAX_DISTINCT"]
        self.assertLessEqual(self.col["distinct"], cap + 1)
        if DISTINCT > cap:
            self.assertTrue(self.col.get("distinct_capped"), "acima do tecto, o registo di-lo")

    def test_counts_stay_true_even_when_samples_are_capped(self):
        """O tecto limita o que se guarda, nunca o que se conta."""
        self.assertEqual(self.col["rows_nonempty"], DISTINCT)

    def test_untruncated_column_declares_column_scope(self):
        self.assertEqual(self.col["stats_scope"], "column")


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl not installed")
class TestEmptyGridIsDroppedNotTruncated(unittest.TestCase):
    """Uma folha com a grelha formatada até ao fim declara linhas que não têm nada. Largá-las não é
    truncar — no workbook real, uma folha declarava 63 461 linhas para NOVE com valores."""

    @classmethod
    def setUpClass(cls):
        from openpyxl.styles import Font
        cls.tmp = Path(tempfile.mkdtemp(prefix="scan-grid-"))
        cls.path = cls.tmp / "grelha.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Grelha"
        ws["A1"], ws["B1"] = "Chave", "Valor"
        for r in range(2, 12):
            ws[f"A{r}"], ws[f"B{r}"] = f"K{r}", r * 3
        for r in (FAR_ROW // 2, FAR_ROW):        # células formatadas, sem valor
            ws.cell(row=r, column=1).font = Font(bold=True)
        wb.save(cls.path)
        cls.ex = extract(cls.path)
        cls.sheet = cls.ex["sheets"][0]

    def test_content_end_ignores_formatted_empty_cells(self):
        self.assertEqual(self.sheet["data_end_content"], 11,
                         f"conteúdo acaba na linha 11; declarado {self.sheet['data_end_declared']}")
        self.assertGreaterEqual(self.sheet["data_end_declared"], FAR_ROW)

    def test_dropping_an_empty_grid_is_not_a_truncation(self):
        self.assertFalse(self.sheet.get("scan_truncated"),
                         "largar grelha vazia não é cortar conteúdo — `scan_truncated` mentiria")
        self.assertEqual(self.sheet["data_end_row"], 11)

    def test_stats_of_a_dropped_grid_still_describe_the_column(self):
        col = next(c for c in self.sheet["columns"] if c["column"] == "B")
        self.assertEqual(col["stats_scope"], "column")
        self.assertEqual(col["max"], 33)

    def test_content_rows_is_published_for_the_whole_workbook(self):
        self.assertIn("content_rows", self.ex["workbook"])
        self.assertEqual(self.ex["workbook"]["content_rows"]["Grelha"], 11)


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl not installed")
class TestSampleScopedStats(unittest.TestCase):
    """Numa folha truncada, `min`/`max`/`distinct` descrevem a amostra varrida, nunca a coluna:
    um `max_date` lido das primeiras linhas de uma série passaria por fim da série."""

    def test_truncated_sheet_marks_its_stats_as_sample(self):
        tmp = Path(tempfile.mkdtemp(prefix="scan-scope-"))
        path = tmp / "longa.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"], ws["B1"] = "Chave", "Valor"
        for r in range(2, CAP + 200):          # ultrapassa o tecto com dados verdadeiros
            ws[f"A{r}"] = f"K{r:06d}"
            ws[f"B{r}"] = r
        wb.save(path)
        sheet = extract(path)["sheets"][0]
        self.assertTrue(sheet["scan_truncated"])
        col = next(c for c in sheet["columns"] if c["column"] == "B")
        self.assertEqual(col["stats_scope"], "sample")
        self.assertLess(col["max"], CAP + 199, "o máximo é o da amostra, e o registo não o dá por da coluna")
        self.assertLessEqual(col["rows_nonempty"], CAP)


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl not installed")
class TestMemoryFailureDegrades(unittest.TestCase):
    def test_memory_error_writes_the_artefact_and_returns_zero(self):
        """A falha da falha (3b-6): no livro de 37 MB o segundo `MemoryError`, levantado durante o
        tratamento do primeiro, matava o processo com rc=1 **depois** de o artefacto estar escrito."""
        self.assertIn("_fail_extraction", X, "Fase 3c: o tratamento da falha é função própria, logo testável")
        tmp = Path(tempfile.mkdtemp(prefix="scan-mem-"))
        out = tmp / "x.extraction.json"
        # um doc já povoado, como está quando a memória acaba a meio do varrimento
        doc = {"artefact": "aisa.capture.extraction", "identity": {"filename": "estoira.xlsx"},
               "sheets": [{"name": "grande", "columns": [{"column": "A"}] * 500}],
               "workbook": {"named_ranges": [{"name": f"N{i}"} for i in range(1000)]}}
        rc = X["_fail_extraction"](doc, str(out), "estoira.xlsx", MemoryError(), None)
        self.assertEqual(rc, 0, "falha de memória degrada com artefacto, não mata o processo")
        self.assertTrue(out.exists(), "o artefacto de falha tem de ficar escrito")
        written = json.loads(out.read_text(encoding="utf-8"))
        self.assertEqual(written["status"], "failed")
        self.assertIn("MemoryError", written["reason"])
        self.assertIn("capability_boundary", written, "mesmo na falha, a fronteira declara o que se sabe")
        self.assertNotIn("sheets", written, "as referências grandes são largadas antes de escrever")

    def test_failure_handler_survives_an_unwritable_target(self):
        """Se nem o artefacto se consegue escrever, o processo continua a sair com 0.

        O alvo tem de ser mesmo inescrevivel: `_write_json` cria os directorios, portanto um
        caminho tipo `Z:/...` e so um caminho relativo em POSIX -- escrevia-se com exito, na
        raiz do repo, e o teste passava sem nunca tocar no ramo que diz testar. Um pai que e
        ficheiro regular falha com ENOTDIR em POSIX e em Windows, sem depender de permissoes.
        """
        tmp = Path(tempfile.mkdtemp(prefix="scan-unwritable-"))
        barreira = tmp / "sou-um-ficheiro"
        barreira.write_text("x", encoding="utf-8")
        alvo = barreira / "sem-permissao" / "x.json"
        doc = {"identity": {"filename": "x.xlsx"}}
        rc = X["_fail_extraction"](doc, str(alvo), "x.xlsx", MemoryError(), None)
        self.assertEqual(rc, 0)
        self.assertFalse(alvo.exists(), "o alvo era inescrevivel: nada pode ter ficado escrito")
        self.assertEqual([p.name for p in tmp.iterdir()], ["sou-um-ficheiro"],
                         "a degradacao nao inventa caminhos alternativos")
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    unittest.main(verbosity=1)
