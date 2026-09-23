"""Captura de cálculo — Fase 3 (P7, P8): severidade por forma, rótulos não são overrides.

Escrito **vermelho** na Fase 1. O que fixa, medido no engagement real:

  - `Inputs!F1319 = =F1316` — a sobreposição manual da cotação do próprio dia (carry-forward
    de 6.ª feira), que é a regra de negócio central da 2.ª feira — saiu `medium`, ao lado de
    `H13='USD'`. Passa a `high`: referência à própria coluna, noutra linha, na última linha
    povoada de uma coluna alimentada por feed.
  - 57 dos 114 pattern exceptions eram texto de cabeçalho em bandas de rótulo dentro do
    intervalo de dados, classificado como «manual override of automated logic». Um rótulo
    não é override. Mas um texto **único** numa coluna calculada (`DGEG ISENTO` no real)
    continua a ser override de negócio e continua `medium` — a fixture planta um.

Na fixture: última linha povoada = `Inputs!F37` (2026-08-10) com `=F34`; `F18 = =F17+10`;
`F12 = 1234.5` (valor tipado); banda de rótulos na linha 13 de `Outputs` (`W13='Produto C'`).

    python .claude/tests/test_replay_severity.py
"""
from __future__ import annotations

import json
import runpy
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.dont_write_bytecode = True
X = runpy.run_path(str(ROOT / "library" / "kernel" / "tools" / "xlsx_extract.py"))
FIXTURE = ROOT / ".claude" / "tests" / "fixtures" / "calc-capture" / "fx-calc-bancas" / "inputs" / "precos-bancas-reduzido.xlsx"
MONDAY_ROW = 37          # 2026-08-10 — última linha povoada de Inputs
PLUS10_ROW = 18          # 2026-07-22 — `=F17+10`
TYPED_ROW = 12           # 2026-07-16 — 1234.5 tipado

try:
    from openpyxl import load_workbook
    HAVE_OPENPYXL = True
except Exception:  # pragma: no cover
    HAVE_OPENPYXL = False


def _replay():
    tmp = Path(tempfile.mkdtemp(prefix="calc-sev-"))
    out = tmp / "x.extraction.json"
    assert X["extract"](str(FIXTURE), str(out), True, None) == 0
    extraction = json.loads(out.read_text(encoding="utf-8"))
    rep = X["Replayer"](load_workbook(str(FIXTURE), data_only=True), extraction)
    rep._fwb = load_workbook(str(FIXTURE), data_only=False)
    rep.run()
    return rep


def _pe(rep, needle):
    return [f for f in rep.findings if f.check == "pattern exceptions" and needle in f.location]


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl not installed")
class TestSeverityByShape(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.rep = _replay()

    def test_live_data_override_on_last_populated_row_is_high(self):
        hits = _pe(self.rep, f"Inputs!F{MONDAY_ROW}")
        self.assertTrue(hits, f"nenhum finding para Inputs!F{MONDAY_ROW} — a divergência tem de ser reportada célula a célula")
        f = hits[0]
        self.assertEqual(f.severity, "high", f"{f.location}: {f.severity} — {f.found}")
        self.assertIn(f"F{MONDAY_ROW - 3}", f.found, "o achado nomeia a célula de origem do carry-forward")
        self.assertIn("feed", f.found.lower(), "o achado diz que a coluna é alimentada por feed")

    def test_plus_ten_patch_mid_series_is_high_too(self):
        hits = _pe(self.rep, f"Inputs!F{PLUS10_ROW}")
        self.assertTrue(hits, f"Inputs!F{PLUS10_ROW} (=F{PLUS10_ROW - 1}+10) tem de sair como finding próprio")
        self.assertEqual(hits[0].severity, "high", hits[0].found)

    def test_typed_value_override_stays_medium(self):
        hits = [f for f in self.rep.findings if f.check == "pattern exceptions"
                and f.location.startswith("Inputs!F") and "1234.5" in f.found]
        self.assertTrue(hits, "o valor tipado em Inputs!F12 continua a ser reportado")
        self.assertEqual(hits[0].severity, "medium")

    def test_label_band_text_is_not_a_manual_override(self):
        wrong = [f for f in self.rep.findings if f.check == "pattern exceptions"
                 and "manual override" in f.found and ("'Produto C'" in f.found or "'Produto A'" in f.found)]
        self.assertEqual(wrong, [], [f"{f.location}: {f.found}" for f in wrong])
        labels = [f for f in self.rep.findings if f.check == "layout labels"]
        self.assertTrue(labels, "as bandas de rótulo passam a ser reportadas em check própria, `info`")
        self.assertTrue(all(f.severity == "info" for f in labels))

    def test_high_is_in_the_summary_vocabulary(self):
        tmp = Path(tempfile.mkdtemp(prefix="calc-sev-md-")) / "r.md"
        X["render_replay_md"]({"identity": {"filename": FIXTURE.name, "sha256": "0" * 64}}, self.rep, str(tmp))
        md = tmp.read_text(encoding="utf-8")
        row = next(l for l in md.splitlines() if l.startswith("| pattern exceptions |"))
        cells = [c.strip() for c in row.strip("|").split("|")]
        # | check | total | high | medium | low | info |
        self.assertGreater(int(cells[2]), 0, f"a coluna `high` de pattern exceptions continua a zero: {row}")


if __name__ == "__main__":
    unittest.main(verbosity=1)
