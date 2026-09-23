#!/usr/bin/env python3
"""test_text_extract.py — capture-lite acceptance test (aisa Phase B).

Run:  python library/kernel/tools/tests/test_text_extract.py
      (or: python -m unittest discover library/kernel/tools/tests)

Covers exactly what the capture-lite contract promises and nothing more:
docx / pdf / vtt extraction with provenance markers, visible failure,
idempotency (cache-hit + byte-identical re-extraction), the evidence index,
and the Excel regression (capture-lite must not touch the xlsx tier).

Fixtures are tiny and generated in a temp dir — no binary blobs in the repo.
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
TOOLS = os.path.dirname(HERE)
REPO = os.environ.get("AISA_TEST_REPO") or os.path.dirname(
    os.path.dirname(os.path.dirname(TOOLS)))
sys.path.insert(0, TOOLS)

import text_extract  # noqa: E402


# ---------------------------------------------------------------- fixtures

def make_docx(path: str) -> None:
    from docx import Document
    doc = Document()
    doc.add_heading("Current Process", level=1)
    doc.add_paragraph("Pricing is updated every morning by the operations team.")
    doc.add_paragraph("Quotes | with a pipe and a \\ backslash.")
    doc.add_heading("Inputs", level=2)
    doc.add_paragraph("Platts quotations arrive automatically.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "field"
    table.cell(0, 1).text = "source"
    table.cell(1, 0).text = "quote"
    table.cell(1, 1).text = "Platts | Argus"
    doc.save(path)


def _pdf_bytes(page_texts: list[str | None]) -> bytes:
    """Minimal uncompressed PDF. `None` page text => page with no text (image-only case)."""
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    catalog_id = add(b"")           # 1, patched below
    pages_id = add(b"")             # 2, patched below
    font_id = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    kids = []
    for text in page_texts:
        stream = (b"BT /F1 12 Tf 20 150 Td (" + text.encode("ascii") + b") Tj ET\n") if text else b""
        content_id = add(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n"
                         + stream + b"endstream")
        page_id = add(b"<< /Type /Page /Parent " + str(pages_id).encode()
                      + b" 0 R /MediaBox [0 0 200 200] /Contents " + str(content_id).encode()
                      + b" 0 R /Resources << /Font << /F1 " + str(font_id).encode()
                      + b" 0 R >> >> >>")
        kids.append(page_id)

    objects[catalog_id - 1] = b"<< /Type /Catalog /Pages " + str(pages_id).encode() + b" 0 R >>"
    objects[pages_id - 1] = (b"<< /Type /Pages /Kids ["
                             + b" ".join(str(k).encode() + b" 0 R" for k in kids)
                             + b"] /Count " + str(len(kids)).encode() + b" >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += str(i).encode() + b" 0 obj\n" + body + b"\nendobj\n"
    xref_at = len(out)
    out += b"xref\n0 " + str(len(objects) + 1).encode() + b"\n0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (b"trailer\n<< /Size " + str(len(objects) + 1).encode()
            + b" /Root " + str(catalog_id).encode() + b" 0 R >>\nstartxref\n"
            + str(xref_at).encode() + b"\n%%EOF\n")
    return bytes(out)


VTT_FIXTURE = """WEBVTT

NOTE this note block must not become a cue

cue-1
00:00:03.333 --> 00:00:06.963
<v Jorge Estev&#227;o>Comecamos pelo objeto
principal aqui,</v>

cue-2
00:00:06.963 --> 00:00:09.613
<v Jorge Estev&#227;o>que e o ficheiro de Excel.</v>

cue-3
00:00:10.013 --> 00:00:13.453
<v Pedro Ornelas>Sim, eu partilho.</v>

cue-4
00:01:02.000 --> 00:01:05.000
Sem locutor identificado.

cue-5
00:01:06.000 --> 00:01:08.000
Outra fala sem locutor.
"""


def run(*args: str) -> subprocess.CompletedProcess:
    env = dict(os.environ, PYTHONIOENCODING="utf-8")
    return subprocess.run([sys.executable, os.path.join(TOOLS, "text_extract.py"), *args],
                          capture_output=True, text=True, encoding="utf-8", env=env)


def read(path: str) -> str:
    with open(path, encoding="utf-8") as fh:
        return fh.read()


def load(path: str) -> dict:
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


class CaptureLiteBase(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.mkdtemp(prefix="aisa-capture-lite-")
        self.inputs = os.path.join(self.tmp, "inputs")
        self.capture = os.path.join(self.tmp, "_capture")
        os.makedirs(self.inputs)
        os.makedirs(self.capture)

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)

    def out(self, name: str) -> tuple[str, str]:
        return (os.path.join(self.capture, f"{name}.text.md"),
                os.path.join(self.capture, f"{name}.extraction.json"))

    def extract(self, name: str, *extra: str) -> subprocess.CompletedProcess:
        md, js = self.out(name)
        return run(os.path.join(self.inputs, name), md, js,
                   "--log", os.path.join(self.capture, "_capture-log.md"), *extra)


# ---------------------------------------------------------------- docx

class TestDocx(CaptureLiteBase):
    def setUp(self) -> None:
        super().setUp()
        self.name = "requirements.docx"
        make_docx(os.path.join(self.inputs, self.name))

    def test_extraction_is_readable_and_traceable(self) -> None:
        proc = self.extract(self.name)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        md, js = self.out(self.name)
        text = read(md)
        doc = load(js)

        self.assertEqual(doc["status"], "ok")
        self.assertEqual(doc["identity"]["format"], ".docx")
        self.assertEqual(doc["provenance"]["cite_as"], "<file> · §<heading> ¶NN")
        self.assertEqual(doc["units"], {"unit": "paragraph", "paragraphs": 3,
                                        "headings": 2, "tables": 1})

        self.assertIn(f"# Source: {self.name}", text)
        self.assertIn("## Current Process", text)          # heading hierarchy preserved
        self.assertIn("### Inputs", text)                  # level 2 stays below level 1
        self.assertIn("[¶1] Pricing is updated every morning", text)
        self.assertIn("[¶3] Platts quotations arrive automatically.", text)
        self.assertIn("[table 1]", text)
        self.assertIn("| field | source |", text)
        self.assertIn(r"| quote | Platts \| Argus |", text)   # cell pipe escaped, not dropped
        self.assertIn("Quotes | with a pipe", text)        # body text kept verbatim

    def test_paragraph_order_preserved(self) -> None:
        self.extract(self.name)
        text = read(self.out(self.name)[0])
        self.assertLess(text.index("## Current Process"), text.index("[¶1]"))
        self.assertLess(text.index("[¶1]"), text.index("[¶2]"))
        self.assertLess(text.index("### Inputs"), text.index("[¶3]"))
        self.assertLess(text.index("[¶3]"), text.index("[table 1]"))

    def test_no_semantic_interpretation(self) -> None:
        self.extract(self.name)
        text = read(self.out(self.name)[0]).lower()
        for banned in ("summary", "requirement:", "risk:", "recommend", "finding"):
            self.assertNotIn(banned, text)


# ---------------------------------------------------------------- pdf

class TestPdf(CaptureLiteBase):
    def test_page_aware_extraction(self) -> None:
        name = "process.pdf"
        with open(os.path.join(self.inputs, name), "wb") as fh:
            fh.write(_pdf_bytes(["PAGE ONE ALPHA", "PAGE TWO BETA"]))
        proc = self.extract(name)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        text = read(self.out(name)[0])
        doc = load(self.out(name)[1])

        self.assertEqual(doc["status"], "ok")
        self.assertEqual(doc["units"]["pages"], 2)
        self.assertEqual(doc["units"]["pages_without_text"], 0)
        self.assertEqual(doc["provenance"]["cite_as"], "<file> · p.N")
        self.assertIn("## [p.1]", text)
        self.assertIn("## [p.2]", text)
        self.assertIn("PAGE ONE ALPHA", text)
        self.assertIn("PAGE TWO BETA", text)
        self.assertLess(text.index("## [p.1]"), text.index("## [p.2]"))
        self.assertLess(text.index("PAGE ONE ALPHA"), text.index("## [p.2]"))

    def test_image_only_pdf_fails_visibly(self) -> None:
        name = "scan.pdf"
        raw = os.path.join(self.inputs, name)
        with open(raw, "wb") as fh:
            fh.write(_pdf_bytes([None, None]))
        proc = self.extract(name)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        doc = load(self.out(name)[1])
        self.assertEqual(doc["status"], "failed")
        self.assertIn("no extractable text", doc["reason"])
        self.assertIn("no OCR", doc["reason"])
        self.assertIn("status: **failed**", read(self.out(name)[0]))
        self.assertTrue(os.path.exists(raw))               # raw source preserved


# ---------------------------------------------------------------- vtt

class TestVtt(CaptureLiteBase):
    def setUp(self) -> None:
        super().setUp()
        self.name = "kickoff.vtt"
        with open(os.path.join(self.inputs, self.name), "w", encoding="utf-8") as fh:
            fh.write(VTT_FIXTURE)

    def test_timestamp_and_speaker_aware(self) -> None:
        proc = self.extract(self.name)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        text = read(self.out(self.name)[0])
        doc = load(self.out(self.name)[1])

        self.assertEqual(doc["status"], "ok")
        self.assertEqual(doc["units"]["cues"], 5)
        self.assertEqual(doc["units"]["passages"], 4)      # cue-1 + cue-2 merged
        self.assertEqual(doc["units"]["speaker_names"], ["Jorge Estevão", "Pedro Ornelas"])
        self.assertEqual(doc["provenance"]["cite_as"], "<file> · [HH:MM:SS] <speaker>")

        self.assertIn("[00:00:03–00:00:09] **Jorge Estevão:**", text)   # entity decoded, run merged
        self.assertIn("Comecamos pelo objeto principal aqui, que e o ficheiro de Excel.", text)
        self.assertIn("[00:00:10–00:00:13] **Pedro Ornelas:** Sim, eu partilho.", text)
        self.assertIn("[00:01:02–00:01:05] Sem locutor identificado.", text)
        self.assertIn("[00:01:06–00:01:08] Outra fala sem locutor.", text)

    def test_transport_noise_dropped_evidence_kept(self) -> None:
        self.extract(self.name)
        text = read(self.out(self.name)[0])
        body = text[text.index("---"):]
        self.assertNotIn("WEBVTT", body)
        self.assertNotIn("-->", body)
        self.assertNotIn("<v ", body)
        self.assertNotIn("cue-1", body)
        self.assertNotIn("this note block", body)          # NOTE block is not a cue

    def test_unattributed_cues_are_not_merged(self) -> None:
        self.extract(self.name)
        text = read(self.out(self.name)[0])
        self.assertIn("Sem locutor identificado.", text)
        self.assertNotIn("Sem locutor identificado. Outra fala", text)

    def test_empty_transcript_is_stated_not_omitted(self) -> None:
        name = "silent.vtt"
        with open(os.path.join(self.inputs, name), "w", encoding="utf-8") as fh:
            fh.write("WEBVTT\n\nNOTE nothing was said\n")
        self.extract(name)
        doc = load(self.out(name)[1])
        self.assertEqual(doc["status"], "empty")
        self.assertIn("0 cues", doc["reason"])
        self.assertIn("Absence is evidence", read(self.out(name)[0]))


# ---------------------------------------------------------------- failure

class TestFailure(CaptureLiteBase):
    def test_corrupt_source_fails_visibly_and_keeps_raw(self) -> None:
        name = "corrupt.docx"
        raw = os.path.join(self.inputs, name)
        with open(raw, "wb") as fh:
            fh.write(b"this is not a docx, only noise")
        size_before = os.path.getsize(raw)
        proc = self.extract(name)
        self.assertEqual(proc.returncode, 0, proc.stderr)          # continues, does not crash
        md, js = self.out(name)
        doc = load(js)
        self.assertEqual(doc["status"], "failed")
        self.assertTrue(doc["reason"])
        self.assertIn("FAILED", proc.stdout)
        self.assertIn("status: **failed**", read(md))
        self.assertIn("Read the raw source", read(md))
        self.assertEqual(os.path.getsize(raw), size_before)        # raw untouched

    def test_failure_is_not_an_epistemic_state(self) -> None:
        name = "corrupt.docx"
        with open(os.path.join(self.inputs, name), "wb") as fh:
            fh.write(b"noise")
        self.extract(name)
        text = read(self.out(name)[0])
        for state in ("Unknown", "Assumed", "Confirmed", "Conflicted", "Risky"):
            self.assertNotIn(state, text)

    def test_unsupported_format_is_a_usage_error_not_an_artefact(self) -> None:
        name = "deck.pptx"
        with open(os.path.join(self.inputs, name), "wb") as fh:
            fh.write(b"x")
        proc = self.extract(name)
        self.assertEqual(proc.returncode, 2)
        self.assertIn("unsupported format", proc.stderr)
        self.assertFalse(os.path.exists(self.out(name)[0]))

    def test_missing_input_is_a_usage_error(self) -> None:
        proc = self.extract("absent.vtt")
        self.assertEqual(proc.returncode, 2)
        self.assertIn("input not found", proc.stderr)


# ---------------------------------------------------------------- idempotency

class TestIdempotency(CaptureLiteBase):
    def setUp(self) -> None:
        super().setUp()
        self.name = "notes.vtt"
        with open(os.path.join(self.inputs, self.name), "w", encoding="utf-8") as fh:
            fh.write(VTT_FIXTURE)

    def test_second_run_is_a_cache_hit(self) -> None:
        self.extract(self.name)
        md, js = self.out(self.name)
        first_md, first_json = read(md), read(js)
        stamp = os.path.getmtime(md)

        proc = self.extract(self.name)
        self.assertIn("cache-hit", proc.stdout)
        self.assertEqual(read(md), first_md)
        self.assertEqual(read(js), first_json)
        self.assertEqual(os.path.getmtime(md), stamp)          # not even rewritten

    def test_forced_re_extraction_is_byte_identical(self) -> None:
        self.extract(self.name)
        md, js = self.out(self.name)
        first_md, first = read(md), load(js)

        proc = self.extract(self.name, "--force")
        self.assertNotIn("cache-hit", proc.stdout)
        self.assertEqual(read(md), first_md)                   # byte-identical extraction

        second = load(js)
        for doc in (first, second):
            doc.pop("extracted_at")                            # the only non-deterministic field
        self.assertEqual(first, second)

    def test_changed_source_re_extracts(self) -> None:
        self.extract(self.name)
        path = os.path.join(self.inputs, self.name)
        with open(path, "a", encoding="utf-8") as fh:
            fh.write("\ncue-6\n00:02:00.000 --> 00:02:04.000\nUma fala nova.\n")
        proc = self.extract(self.name)
        self.assertNotIn("cache-hit", proc.stdout)
        self.assertIn("Uma fala nova.", read(self.out(self.name)[0]))

    def test_no_duplicate_artefacts_accumulate(self) -> None:
        for _ in range(3):
            self.extract(self.name)
        produced = sorted(n for n in os.listdir(self.capture) if n.startswith(self.name))
        self.assertEqual(produced, [f"{self.name}.extraction.json", f"{self.name}.text.md"])


# ---------------------------------------------------------------- evidence index

class TestEvidenceIndex(CaptureLiteBase):
    def setUp(self) -> None:
        super().setUp()
        make_docx(os.path.join(self.inputs, "requirements.docx"))
        with open(os.path.join(self.inputs, "kickoff.vtt"), "w", encoding="utf-8") as fh:
            fh.write(VTT_FIXTURE)
        with open(os.path.join(self.inputs, "scan.pdf"), "wb") as fh:
            fh.write(_pdf_bytes([None]))
        with open(os.path.join(self.inputs, "deck.pptx"), "wb") as fh:
            fh.write(b"x")
        for name in ("requirements.docx", "kickoff.vtt", "scan.pdf"):
            self.extract(name)

    def test_index_lists_every_source_with_its_status(self) -> None:
        proc = run("--index", self.tmp)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        index = read(os.path.join(self.capture, "evidence-index.md"))
        self.assertIn("| `requirements.docx` | `.docx` | `requirements.docx.text.md` | `ok` |", index)
        self.assertIn("| `kickoff.vtt` | `.vtt` | `kickoff.vtt.text.md` | `ok` |", index)
        self.assertIn("| `scan.pdf` | `.pdf` | `scan.pdf.text.md` | `failed` |", index)
        self.assertIn("| `deck.pptx` | `.pptx` | — | `not captured` |", index)
        self.assertIn("authoritative on conflict", index)

    def test_index_regeneration_is_idempotent(self) -> None:
        run("--index", self.tmp)
        first = read(os.path.join(self.capture, "evidence-index.md"))
        run("--index", self.tmp)
        self.assertEqual(read(os.path.join(self.capture, "evidence-index.md")), first)


# ---------------------------------------------------------------- excel regression

XLSX_ENGAGEMENT = os.path.join(REPO, "projects", "pricing-marinha")


class TestExcelRegression(unittest.TestCase):
    """capture-lite must not touch, replace or shadow the xlsx tier."""

    def test_capture_lite_refuses_xlsx(self) -> None:
        self.assertNotIn(".xlsx", text_extract.EXTRACTORS)
        self.assertNotIn(".xlsm", text_extract.EXTRACTORS)
        self.assertEqual(text_extract.XLSX_FORMATS, (".xlsx", ".xlsm"))

    def test_xlsx_artefact_id_is_distinct(self) -> None:
        self.assertNotEqual(text_extract.ARTEFACT_ID, text_extract.XLSX_ARTEFACT_ID)
        self.assertEqual(text_extract.XLSX_ARTEFACT_ID, "aisa.capture.extraction")

    @unittest.skipUnless(os.path.isdir(XLSX_ENGAGEMENT), "pricing-marinha not present")
    def test_existing_xlsx_extraction_still_re_extracts_identically(self) -> None:
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("openpyxl not installed")
        src = None
        for name in os.listdir(os.path.join(XLSX_ENGAGEMENT, "inputs")):
            if name.lower().endswith((".xlsx", ".xlsm")):
                src = name
                break
        if not src:
            self.skipTest("no xlsx input in pricing-marinha")
        committed_path = os.path.join(XLSX_ENGAGEMENT, "_capture", f"{src}.extraction.json")
        if not os.path.exists(committed_path):
            self.skipTest("no committed extraction to compare against")
        committed = load(committed_path)

        tmp = tempfile.mkdtemp(prefix="aisa-xlsx-regression-")
        try:
            out = os.path.join(tmp, f"{src}.extraction.json")
            env = dict(os.environ, PYTHONIOENCODING="utf-8")
            proc = subprocess.run(
                [sys.executable, os.path.join(TOOLS, "xlsx_extract.py"),
                 os.path.join(XLSX_ENGAGEMENT, "inputs", src), out],
                capture_output=True, text=True, encoding="utf-8", env=env)
            self.assertEqual(proc.returncode, 0, proc.stderr)
            fresh = load(out)
            self.assertEqual(fresh["status"], "ok")
            self.assertEqual(fresh["identity"]["sha256"], committed["identity"]["sha256"])
            self.assertEqual(fresh["sheets"], committed["sheets"])
            self.assertEqual(fresh["workbook"], committed["workbook"])

            replay_out = os.path.join(tmp, f"{src}.replay.md")
            proc = subprocess.run(
                [sys.executable, os.path.join(TOOLS, "xlsx_extract.py"), "--replay",
                 os.path.join(XLSX_ENGAGEMENT, "inputs", src), out, replay_out],
                capture_output=True, text=True, encoding="utf-8", env=env)
            self.assertEqual(proc.returncode, 0, proc.stderr)
            committed_replay = os.path.join(XLSX_ENGAGEMENT, "_capture", f"{src}.replay.md")
            if os.path.exists(committed_replay):
                def findings(md: str) -> str:
                    # drop what is relative to the run date, not to the
                    # workbook: replay timestamp + staleness aging rows
                    return "\n".join(ln for ln in md.splitlines()
                                     if "Replayed:" not in ln
                                     and "distribution {" not in ln
                                     and "replay date" not in ln)
                self.assertEqual(findings(read(replay_out)),
                                 findings(read(committed_replay)))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    @unittest.skipUnless(os.path.isdir(XLSX_ENGAGEMENT), "pricing-marinha not present")
    def test_index_reports_the_xlsx_tier_not_the_text_tier(self) -> None:
        tmp = tempfile.mkdtemp(prefix="aisa-index-xlsx-")
        try:
            os.makedirs(os.path.join(tmp, "inputs"))
            os.makedirs(os.path.join(tmp, "_capture"))
            src = next(n for n in os.listdir(os.path.join(XLSX_ENGAGEMENT, "inputs"))
                       if n.lower().endswith((".xlsx", ".xlsm")))
            with open(os.path.join(tmp, "inputs", src), "wb") as fh:
                fh.write(b"placeholder")
            shutil.copy(os.path.join(XLSX_ENGAGEMENT, "_capture", f"{src}.extraction.json"),
                        os.path.join(tmp, "_capture", f"{src}.extraction.json"))
            run("--index", tmp)
            index = read(os.path.join(tmp, "_capture", "evidence-index.md"))
            self.assertIn("`process-model.md` (PM ids)", index)
            self.assertIn("PM-NNN → <sheet>!<cell>", index)
            self.assertNotIn(f"{src}.text.md", index)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    unittest.main(verbosity=2)
